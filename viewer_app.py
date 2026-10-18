# -*- coding: utf-8 -*-
"""
商户数据查看器
- 一行选择：省 → 市 → 区县 → 行业 → 细分类目（同时驱动表格筛选和采集）
- 数据范围：全部数据 / 仅新增未导出 / 仅已导出（已导出的也能看）
- Excel导出：分割（每份条数/分成份数）+ 电话去重（任何两次导出零重复）
- 无任何使用限制
"""
import tkinter as tk
from tkinter import ttk, filedialog, messagebox
import requests
import threading
import json
import os
import re
import sys
import time
import shutil
import hashlib
import webbrowser
from datetime import datetime

# ============ 配置 ============
# 请填入你自己的高德开放平台 Web服务 API Key（https://lbs.amap.com/ 免费申请）
# 本地打包版由 build_exe.py 自动从 key.txt 注入，仓库代码永远不含真实 Key
AMAP_KEY = ''

HISTORY_FILE = '已导出电话记录.json'
HISTORY_BAK = HISTORY_FILE + '.bak'
CACHE_FILE = '采集数据缓存.json'
EXPORT_LOG_FILE = '导出日志.json'

# 全国省级行政区（34个，静态内置，避免启动时请求）
PROVINCES = [
    ('110000', '北京市'), ('120000', '天津市'), ('130000', '河北省'), ('140000', '山西省'),
    ('150000', '内蒙古自治区'), ('210000', '辽宁省'), ('220000', '吉林省'), ('230000', '黑龙江省'),
    ('310000', '上海市'), ('320000', '江苏省'), ('330000', '浙江省'), ('340000', '安徽省'),
    ('350000', '福建省'), ('360000', '江西省'), ('370000', '山东省'), ('410000', '河南省'),
    ('420000', '湖北省'), ('430000', '湖南省'), ('440000', '广东省'), ('450000', '广西壮族自治区'),
    ('460000', '海南省'), ('500000', '重庆市'), ('510000', '四川省'), ('520000', '贵州省'),
    ('530000', '云南省'), ('540000', '西藏自治区'), ('610000', '陕西省'), ('620000', '甘肃省'),
    ('630000', '青海省'), ('640000', '宁夏回族自治区'), ('650000', '新疆维吾尔自治区'),
    ('710000', '台湾省'), ('810000', '香港特别行政区'), ('820000', '澳门特别行政区'),
]

# 宜昌13区县（离线兜底：行政区API取不到时用）
YICHANG_DISTRICTS = [
    ('420502', '西陵区'), ('420503', '伍家岗区'), ('420504', '点军区'), ('420505', '猇亭区'),
    ('420506', '夷陵区'), ('420525', '远安县'), ('420526', '兴山县'), ('420527', '秭归县'),
    ('420528', '长阳土家族自治县'), ('420529', '五峰土家族自治县'), ('420581', '宜都市'),
    ('420582', '当阳市'), ('420583', '枝江市'),
]

# ============ 行业树：一级 -> 二级 -> 查询组 ============
_FOOD_SUBS = [
    ('中餐厅', [
        ('types', '050100', '中餐厅'), ('types', '050101', '中餐厅'), ('types', '050102', '中餐厅'),
        ('types', '050103', '中餐厅'), ('types', '050104', '中餐厅'), ('types', '050105', '中餐厅'),
        ('types', '050106', '中餐厅'), ('types', '050107', '中餐厅'), ('types', '050108', '中餐厅'),
        ('types', '050109', '中餐厅'), ('types', '050110', '中餐厅'), ('types', '050111', '中餐厅'),
        ('types', '050112', '中餐厅'), ('types', '050113', '中餐厅'), ('types', '050114', '中餐厅'),
        ('types', '050115', '中餐厅'), ('types', '050116', '中餐厅'), ('types', '050117', '中餐厅'),
        ('types', '050118', '中餐厅'), ('types', '050119', '中餐厅'), ('types', '050120', '中餐厅'),
        ('types', '050121', '中餐厅'), ('types', '050122', '中餐厅'), ('types', '050123', '中餐厅'),
    ]),
    ('外国餐厅', [
        ('types', '050200', '外国餐厅'), ('types', '050201', '外国餐厅'), ('types', '050202', '外国餐厅'),
        ('types', '050203', '外国餐厅'), ('types', '050204', '外国餐厅'), ('types', '050205', '外国餐厅'),
        ('types', '050206', '外国餐厅'), ('types', '050207', '外国餐厅'), ('types', '050208', '外国餐厅'),
        ('types', '050209', '外国餐厅'), ('types', '050210', '外国餐厅'), ('types', '050211', '外国餐厅'),
        ('types', '050212', '外国餐厅'), ('types', '050213', '外国餐厅'),
    ]),
    ('快餐厅', [
        ('types', '050300', '快餐厅'), ('types', '050301', '快餐厅'), ('types', '050302', '快餐厅'),
        ('types', '050303', '快餐厅'), ('types', '050304', '快餐厅'), ('types', '050305', '快餐厅'),
        ('types', '050306', '快餐厅'), ('types', '050307', '快餐厅'), ('types', '050308', '快餐厅'),
        ('types', '050309', '快餐厅'), ('types', '050310', '快餐厅'), ('types', '050311', '快餐厅'),
    ]),
    ('休闲餐饮', [('types', '050400', '休闲餐饮')]),
    ('咖啡厅', [
        ('types', '050500', '咖啡厅'), ('types', '050501', '咖啡厅'), ('types', '050502', '咖啡厅'),
        ('types', '050503', '咖啡厅'), ('types', '050504', '咖啡厅'),
    ]),
    ('茶艺馆', [('types', '050600', '茶艺馆')]),
    ('冷饮店', [('types', '050700', '冷饮店')]),
    ('糕饼店', [('types', '050800', '糕饼店')]),
    ('甜品店', [('types', '050900', '甜品店')]),
    ('其他餐饮', [('types', '050000', '其他餐饮')]),
]

_SHOPPING_SUBS = [
    ('综合商场', [('types', '060100', '综合商场'), ('kw', '购物中心', '综合商场'), ('kw', '百货商场', '综合商场')]),
    ('超市', [('types', '060400', '超市'), ('kw', '超市', '超市')]),
    ('便利店', [('types', '060200', '便利店')]),
    ('市场集市', [('types', '060700', '市场集市'), ('kw', '批发市场', '市场集市'), ('kw', '农贸市场', '市场集市'), ('kw', '菜市场', '市场集市')]),
    ('家电数码', [('types', '060300', '家电数码'), ('kw', '手机卖场', '家电数码'), ('kw', '电脑城', '家电数码')]),
    ('建材家居', [('kw', '建材市场', '建材家居'), ('kw', '家居广场', '建材家居'), ('kw', '家具城', '建材家居')]),
    ('服装鞋帽', [('kw', '服装店', '服装鞋帽'), ('kw', '鞋店', '服装鞋帽')]),
    ('眼镜珠宝书店', [('kw', '眼镜店', '眼镜珠宝书店'), ('kw', '珠宝店', '眼镜珠宝书店'), ('kw', '书店', '眼镜珠宝书店')]),
    ('花鸟宠物', [('kw', '花鸟市场', '花鸟宠物'), ('kw', '宠物店', '花鸟宠物')]),
]

_LIFE_SUBS = [
    ('美容美发', [('kw', '美发', '美容美发'), ('kw', '理发店', '美容美发'), ('kw', '美容院', '美容美发')]),
    ('美甲美睫', [('kw', '美甲', '美甲美睫'), ('kw', '美睫', '美甲美睫')]),
    ('洗浴足浴按摩', [('kw', '洗浴', '洗浴足浴按摩'), ('kw', '足浴', '洗浴足浴按摩'), ('kw', '足疗', '洗浴足浴按摩'), ('kw', '按摩推拿', '洗浴足浴按摩')]),
    ('洗衣干洗', [('kw', '洗衣店', '洗衣干洗'), ('kw', '干洗店', '洗衣干洗')]),
    ('摄影照相', [('kw', '摄影', '摄影照相'), ('kw', '照相馆', '摄影照相'), ('kw', '婚纱摄影', '摄影照相')]),
    ('家政保姆', [('kw', '家政', '家政保姆'), ('kw', '保姆月嫂', '家政保姆')]),
    ('家电手机维修', [('kw', '家电维修', '家电手机维修'), ('kw', '手机维修', '家电手机维修')]),
    ('物流快递', [('types', '070700', '物流快递'), ('kw', '快递网点', '物流快递'), ('kw', '菜鸟驿站', '物流快递')]),
    ('殡葬服务', [('types', '070900', '殡葬服务')]),
    ('营业厅', [('kw', '移动营业厅', '营业厅'), ('kw', '联通营业厅', '营业厅'), ('kw', '电信营业厅', '营业厅')]),
]

_HOTEL_SUBS = [
    ('宾馆酒店', [('types', '100100', '宾馆酒店'), ('kw', '酒店', '宾馆酒店'), ('kw', '宾馆', '宾馆酒店')]),
    ('民宿客栈', [('kw', '民宿', '民宿客栈'), ('kw', '客栈', '民宿客栈')]),
    ('旅馆招待所', [('types', '100200', '旅馆招待所'), ('kw', '旅馆', '旅馆招待所'), ('kw', '招待所', '旅馆招待所')]),
    ('公寓住宿', [('kw', '酒店式公寓', '公寓住宿'), ('kw', '电竞酒店', '公寓住宿'), ('kw', '度假酒店', '公寓住宿')]),
]

_AUTO_SUBS = [
    ('汽车销售4S店', [('types', '020000', '汽车销售4S店'), ('kw', '汽车4S店', '汽车销售4S店'), ('kw', '汽车销售', '汽车销售4S店')]),
    ('汽车维修保养', [('types', '030000', '汽车维修保养'), ('kw', '汽车维修', '汽车维修保养'), ('kw', '汽车保养', '汽车维修保养'), ('kw', '汽修厂', '汽车维修保养')]),
    ('加油加气充电', [('types', '010100', '加油加气充电'), ('types', '010200', '加油加气充电'), ('types', '011100', '加油加气充电')]),
    ('洗车美容', [('types', '010400', '洗车美容'), ('kw', '洗车', '洗车美容'), ('kw', '汽车美容', '洗车美容')]),
    ('驾校检测站', [('kw', '驾校', '驾校检测站'), ('kw', '汽车检测', '驾校检测站'), ('kw', '机动车年检', '驾校检测站')]),
    ('摩托车电动车', [('types', '040000', '摩托车电动车'), ('kw', '电动车行', '摩托车电动车')]),
]

_SPORTS_SUBS = [
    ('KTV夜总会', [('types', '080302', 'KTV夜总会'), ('kw', 'KTV', 'KTV夜总会'), ('kw', '夜总会', 'KTV夜总会')]),
    ('酒吧', [('kw', '酒吧', '酒吧'), ('kw', '清吧', '酒吧'), ('kw', '音乐酒馆', '酒吧')]),
    ('网吧电竞', [('types', '080308', '网吧电竞'), ('kw', '网吧', '网吧电竞'), ('kw', '网咖', '网吧电竞'), ('kw', '电竞馆', '网吧电竞')]),
    ('台球棋牌', [('kw', '台球厅', '台球棋牌'), ('kw', '棋牌室', '台球棋牌'), ('kw', '麻将馆', '台球棋牌')]),
    ('健身游泳', [('kw', '健身房', '健身游泳'), ('kw', '游泳馆', '健身游泳'), ('kw', '瑜伽馆', '健身游泳')]),
    ('电影院', [('types', '080601', '电影院'), ('kw', '电影院', '电影院')]),
    ('游乐场娱乐', [('kw', '游乐场', '游乐场娱乐'), ('kw', '水上乐园', '游乐场娱乐'), ('kw', '密室逃脱', '游乐场娱乐'), ('kw', '剧本杀', '游乐场娱乐')]),
]

_MEDICAL_SUBS = [
    ('综合医院', [('types', '090100', '综合医院'), ('types', '090200', '综合医院'), ('kw', '医院', '综合医院')]),
    ('诊所卫生院', [('types', '090600', '诊所卫生院'), ('kw', '诊所', '诊所卫生院'), ('kw', '卫生院', '诊所卫生院'), ('kw', '卫生室', '诊所卫生院')]),
    ('药店药房', [('kw', '药店', '药店药房'), ('kw', '大药房', '药店药房')]),
    ('口腔眼科中医', [('kw', '口腔诊所', '口腔眼科中医'), ('kw', '牙科', '口腔眼科中医'), ('kw', '眼科', '口腔眼科中医'), ('kw', '中医馆', '口腔眼科中医')]),
    ('宠物医院', [('kw', '宠物医院', '宠物医院'), ('kw', '动物医院', '宠物医院')]),
    ('体检急救', [('kw', '体检中心', '体检急救'), ('kw', '急救中心', '体检急救')]),
]

_EDU_SUBS = [
    ('中小学', [('kw', '中学', '中小学'), ('kw', '小学', '中小学'), ('kw', '学校', '中小学')]),
    ('幼儿园', [('kw', '幼儿园', '幼儿园')]),
    ('培训机构', [('kw', '培训机构', '培训机构'), ('kw', '教育咨询', '培训机构'), ('kw', '托管班', '培训机构')]),
    ('驾校培训', [('kw', '驾校', '驾校培训')]),
    ('大学职校', [('kw', '大学', '大学职校'), ('kw', '职业技术学院', '大学职校')]),
    ('图书馆博物馆', [('kw', '图书馆', '图书馆博物馆'), ('kw', '博物馆', '图书馆博物馆'), ('kw', '文化馆', '图书馆博物馆')]),
]

_FIN_SUBS = [
    ('银行网点', [('types', '160100', '银行网点'), ('kw', '银行', '银行网点')]),
    ('ATM自助银行', [('types', '160300', 'ATM自助银行'), ('kw', 'ATM', 'ATM自助银行')]),
    ('保险证券', [('kw', '保险公司', '保险证券'), ('kw', '证券公司', '保险证券'), ('kw', '营业部', '保险证券')]),
]

_COMPANY_SUBS = [
    ('公司企业', [('types', '170200', '公司企业'), ('kw', '有限公司', '公司企业')]),
    ('工厂园区', [('kw', '工厂', '工厂园区'), ('kw', '工业园', '工厂园区'), ('kw', '产业园', '工厂园区')]),
    ('广告装饰', [('kw', '广告公司', '广告装饰'), ('kw', '装饰工程', '广告装饰'), ('kw', '装修公司', '广告装饰')]),
    ('农林牧渔', [('kw', '养殖', '农林牧渔'), ('kw', '种植合作社', '农林牧渔')]),
]

_SCENIC_SUBS = [
    ('公园广场', [('kw', '公园', '公园广场'), ('kw', '广场', '公园广场')]),
    ('景区景点', [('types', '110200', '景区景点'), ('kw', '景区', '景区景点'), ('kw', '风景区', '景区景点'), ('kw', '旅游景点', '景区景点')]),
    ('寺庙道观', [('kw', '寺庙', '寺庙道观'), ('kw', '道观', '寺庙道观')]),
    ('纪念馆展馆', [('kw', '纪念馆', '纪念馆展馆'), ('kw', '展览馆', '纪念馆展馆')]),
]

_ESTATE_SUBS = [
    ('房产中介', [('types', '120700', '房产中介'), ('kw', '房产中介', '房产中介'), ('kw', '房地产中介', '房产中介'), ('kw', '房屋中介', '房产中介')]),
    ('楼盘售楼处', [('kw', '售楼处', '楼盘售楼处'), ('kw', '营销中心', '楼盘售楼处')]),
    ('物业管理', [('kw', '物业公司', '物业管理')]),
]


def _build_tree():
    raw = [
        ('餐饮服务', _FOOD_SUBS),
        ('购物服务', _SHOPPING_SUBS),
        ('生活服务', _LIFE_SUBS),
        ('住宿服务', _HOTEL_SUBS),
        ('汽车服务', _AUTO_SUBS),
        ('体育休闲服务', _SPORTS_SUBS),
        ('医疗保健服务', _MEDICAL_SUBS),
        ('科教文化服务', _EDU_SUBS),
        ('金融保险服务', _FIN_SUBS),
        ('公司企业', _COMPANY_SUBS),
        ('风景名胜', _SCENIC_SUBS),
        ('房地产中介', _ESTATE_SUBS),
    ]
    tree = {}
    for ind_name, subs in raw:
        d = {'全部': []}
        for sub_name, queries in subs:
            d[sub_name] = queries
            d['全部'].extend(queries)
        tree[ind_name] = d
    return tree


IND_TREE = _build_tree()
IND_NAMES = list(IND_TREE.keys())

HEADERS = ['商户名称', '省份', '城市', '区县', '地址', '联系电话', '经纬度', '类型', '行业', '二级类目']

ALL_DISTRICTS = '全部区县'
SCOPE_ALL = '全部数据'
SCOPE_NEW = '仅新增未导出'
SCOPE_EXPORTED = '仅已导出'


# ============ 工具函数 ============
def safe_str(val):
    if val is None:
        return ''
    if isinstance(val, list):
        return '; '.join(str(v) for v in val if v)
    return str(val).strip()


def split_phones(tel_str):
    if not tel_str:
        return []
    return [p.strip() for p in re.split(r'[;；,，/]', str(tel_str)) if p.strip()]


def normalize_phone(p):
    return re.sub(r'\D', '', str(p or ''))


def industry_of(rec):
    ind = rec.get('industry')
    if ind:
        return ind
    t = (rec.get('type') or '').split(';')
    return t[0].strip() if t and t[0] else ''


def sub_of(rec):
    s = rec.get('sub')
    if s:
        return s
    t = (rec.get('type') or '').split(';')
    return t[1].strip() if len(t) > 1 else ''


def normalize_record(rec):
    # 字段别名统一：内置数据用 province/city/district，采集数据用 pname/cityname/adname
    # 下游（表格/筛选/导出/搜索）统一读 pname/cityname/adname
    if not rec.get('pname') and rec.get('province'):
        rec['pname'] = rec['province']
    if not rec.get('cityname') and rec.get('city'):
        rec['cityname'] = rec['city']
    if not rec.get('adname') and rec.get('district'):
        rec['adname'] = rec['district']
    # 类型兜底：任何来源的数据都洗成字符串，防 tel 是列表导致搜索/显示拼接崩溃
    for k in ('name', 'pname', 'cityname', 'adname', 'address', 'tel', 'location', 'type'):
        rec[k] = safe_str(rec.get(k))
    rec['industry'] = industry_of(rec)
    rec['sub'] = sub_of(rec)
    return rec


def make_chunks(data, mode, value):
    total = len(data)
    if total == 0:
        return []
    if mode == '分成份数':
        if value <= 1 or value >= total:
            return [data]
        per = (total + value - 1) // value
    else:
        per = value
    if per <= 0 or per >= total:
        return [data]
    return [data[i:i + per] for i in range(0, total, per)]


# ============ 高德API（带重试） ============
def amap_get(url, params, retries=2, timeout=15):
    last = None
    for attempt in range(retries + 1):
        try:
            resp = requests.get(url, params=params, timeout=timeout)
            return resp.json()
        except Exception as e:
            last = e
            if attempt < retries:
                time.sleep(0.5 * (attempt + 1))
    raise last


def search_poi(adcode, query, page=1, offset=25):
    kind, value, _label = query
    params = {
        'key': AMAP_KEY,
        'types': value if kind == 'types' else '',
        'keywords': value if kind == 'kw' else '',
        'city': adcode,
        'citylimit': 'true',
        'offset': offset,
        'page': page,
        'extensions': 'all',
        'output': 'json',
    }
    return amap_get('https://restapi.amap.com/v3/place/text', params)


_district_cache = {}


def fetch_sub_districts(adcode):
    """拉取下级行政区。只有请求成功才缓存——失败（网络抖动）不缓存，下次会重试。"""
    if adcode in _district_cache:
        return _district_cache[adcode]
    result = []
    ok = False
    try:
        data = amap_get('https://restapi.amap.com/v3/config/district',
                        {'key': AMAP_KEY, 'keywords': adcode, 'subdistrict': 1, 'extensions': 'base'},
                        retries=1, timeout=8)
        if data.get('status') == '1' and data.get('districts'):
            subs = data['districts'][0].get('districts', [])
            result = [(d.get('adcode', ''), d.get('name', '')) for d in subs if d.get('adcode')]
            ok = True  # 请求成功（哪怕下级确实为空也缓存，那是真实状态）
    except Exception:
        pass
    if ok:
        _district_cache[adcode] = result
    return result


def classify_api_error(data):
    """识别高德API错误类型。返回: 'daily'日额度尽 / 'key' Key无效 / 'qps' 超限可重试 / 'other'"""
    if not isinstance(data, dict):
        return 'other'
    infocode = str(data.get('infocode', ''))
    info = str(data.get('info', '')).upper()
    if infocode == '10003' or 'DAILY_QUERY_OVER_LIMIT' in info:
        return 'daily'
    if infocode in ('10001', '10008', '10009') or 'INVALID_USER_KEY' in info:
        return 'key'
    if infocode in ('10016', '10044', '10019', '10020', '10021') or 'CUQPS' in info or 'QPS' in info:
        return 'qps'
    return 'other'


# ============ 导出 ============
def export_rows(records):
    rows = []
    for r in records:
        rows.append([
            safe_str(r.get('name')), safe_str(r.get('pname')), safe_str(r.get('cityname')),
            safe_str(r.get('adname')), safe_str(r.get('address')), safe_str(r.get('tel')),
            safe_str(r.get('location')), safe_str(r.get('type')),
            safe_str(r.get('industry')), safe_str(r.get('sub')),
        ])
    return rows


def write_excel(filepath, rows):
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = '商户数据'
    header_font = Font(name='微软雅黑', bold=True, color='FFFFFF', size=11)
    header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    for col, h in enumerate(HEADERS, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')
    for r_idx, row in enumerate(rows, 2):
        for c_idx, val in enumerate(row, 1):
            ws.cell(row=r_idx, column=c_idx, value=val).alignment = Alignment(vertical='center')
    widths = {'商户名称': 28, '省份': 8, '城市': 8, '区县': 12, '地址': 40,
              '联系电话': 22, '经纬度': 20, '类型': 24, '行业': 12, '二级类目': 14}
    for col, h in enumerate(HEADERS, 1):
        ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = widths.get(h, 15)
    ws.freeze_panes = 'A2'
    wb.save(filepath)


# ============ 主程序 ============
class MerchantViewer:
    def __init__(self):
        self.root = tk.Tk()
        self.root.title('商户数据查看器')
        self.root.geometry('1500x880')
        self.root.minsize(1150, 620)

        self.records = []
        self.filtered = []
        self.crawling = False
        self.stop_flag = False
        self._district_items = []
        self._city_items = []
        self._region_gen = 0      # 地区级联代际标记：防旧响应覆盖新选择
        self._search_after = None  # 搜索防抖句柄
        self._hist_win = None      # 历史窗口引用：防多开
        self._instance_mutex = None

        if getattr(sys, 'frozen', False):
            self.exe_dir = os.path.dirname(sys.executable)
        else:
            self.exe_dir = os.path.dirname(os.path.abspath(__file__))
        if not self._acquire_single_instance():
            self.root.destroy()
            sys.exit(0)
        # 绿色便携：状态文件放EXE旁边的"数据文件"文件夹，整个文件夹拷走即用，不写C盘
        self._state_dir_fallback = False
        self.state_dir = os.path.join(self.exe_dir, '数据文件')
        if not self._dir_writable(self.exe_dir):
            # EXE所在目录不可写（如 Program Files、只读介质）→ 退回系统目录兜底
            self.state_dir = os.path.join(os.environ.get('LOCALAPPDATA', self.exe_dir),
                                          '商户数据查看器')
            self._state_dir_fallback = True
        else:
            # 对抗场景：EXE旁边已存在一个叫"数据文件"的【文件】（用户乱放）→ 挪开再建目录
            if os.path.isfile(self.state_dir):
                try:
                    os.rename(self.state_dir,
                              self.state_dir + '_冲突备份_' + datetime.now().strftime('%H%M%S'))
                except OSError:
                    pass
        try:
            os.makedirs(self.state_dir, exist_ok=True)
        except OSError:
            # 目录建不起来（权限/占用等）→ 退回系统目录兜底，保证数据有地方存
            self.state_dir = os.path.join(os.environ.get('LOCALAPPDATA', self.exe_dir),
                                          '商户数据查看器')
            self._state_dir_fallback = True
            try:
                os.makedirs(self.state_dir, exist_ok=True)
            except OSError:
                pass
        self.history_path = os.path.join(self.state_dir, HISTORY_FILE)
        self.cache_path = os.path.join(self.state_dir, CACHE_FILE)
        self._migrate_legacy_state()  # 老版本升级：把系统目录里的旧状态搬回EXE旁边

        self._build_ui()
        if self._state_dir_fallback:
            messagebox.showwarning(
                '数据保存位置',
                '软件所在目录不可写，数据文件将临时改存到：\n' + self.state_dir +
                '\n\n建议把软件移到桌面等可写位置后再使用。')
        self._init_history()
        self._load_cache()
        self._refresh_table()

        self.root.protocol('WM_DELETE_WINDOW', self._on_close)
        # 启动地区级联：湖北省 → 宜昌市 → 全部区县
        self.root.after(100, self._init_region_cascade)

    # ---------- 资源路径 ----------
    def resource_path(self, name):
        if getattr(sys, 'frozen', False):
            return os.path.join(sys._MEIPASS, name)
        return os.path.join(self.exe_dir, name)

    @staticmethod
    def _dir_writable(path):
        try:
            test = os.path.join(path, '.write_test_tmp')
            with open(test, 'w') as f:
                f.write('t')
            os.remove(test)
            return True
        except OSError:
            return False

    def _acquire_single_instance(self):
        """同一文件夹只允许开一个实例。双开时两个进程各自读写数据文件，
        后关的覆盖先关的，去重记录会丢——直接拦下第二个。不同文件夹的副本互不影响。"""
        try:
            import ctypes
            # use_last_error=True：保证拿到的是 CreateMutexW 刚设置的错误码，
            # 不会被中间其他调用覆盖（windll 直调拿 GetLastError 在理论上不可靠）
            k32 = ctypes.WinDLL('kernel32', use_last_error=True)
            name = 'Local\\MerchantViewer_' + hashlib.md5(
                os.path.abspath(self.exe_dir).lower().encode('utf-8')).hexdigest()[:16]
            self._instance_mutex = k32.CreateMutexW(None, False, name)
            # ERROR_ALREADY_EXISTS = 183；进程崩溃句柄自动释放，不会误锁
            if ctypes.get_last_error() == 183:
                messagebox.showwarning('已在运行',
                                       '这个软件已经在运行了。\n\n'
                                       '同一个文件夹里只能开一个（双开会互相覆盖数据文件），\n'
                                       '请到任务栏找已在运行的窗口。')
                return False
        except Exception:
            pass  # 拿不到系统接口就放行，正常使用优先
        return True

    # ---------- 老版本状态文件迁移（两个旧位置 → EXE旁边的"数据文件"） ----------
    def _migrate_legacy_state(self):
        """旧版本把 json 存在两处：更老的放EXE同目录，上一版放 %LOCALAPPDATA%。
        绿色化后统一搬进 EXE旁边的"数据文件"文件夹，校验一致后才删原文件。"""
        candidates = [self.exe_dir]  # 更老版本：直接放EXE旁边
        base = os.environ.get('LOCALAPPDATA')
        legacy_dir = None
        if base:
            legacy_dir = os.path.join(base, '商户数据查看器')
            candidates.append(legacy_dir)  # 上一版：系统用户目录
        names = [HISTORY_FILE, HISTORY_BAK, HISTORY_FILE + '.corrupt.bak',
                 CACHE_FILE, EXPORT_LOG_FILE]
        moved = 0
        for src_dir in candidates:
            if not os.path.isdir(src_dir):
                continue
            if os.path.abspath(src_dir) == os.path.abspath(self.state_dir):
                continue  # 兜底场景下 state_dir 可能正好等于系统目录
            for fname in names:
                old_p = os.path.join(src_dir, fname)
                new_p = os.path.join(self.state_dir, fname)
                if not os.path.exists(old_p) or os.path.exists(new_p):
                    continue  # 没有旧文件，或目标已有新状态（新的优先，不覆盖）
                try:
                    shutil.copyfile(old_p, new_p)
                    with open(old_p, 'rb') as f:
                        old_bytes = f.read()
                    with open(new_p, 'rb') as f:
                        new_bytes = f.read()
                    if old_bytes == new_bytes:
                        os.remove(old_p)
                        moved += 1
                except OSError:
                    pass  # 迁移失败就留着，不影响使用
        if moved:
            if legacy_dir:
                try:
                    os.rmdir(legacy_dir)  # 只在空了才能删掉，安全
                except OSError:
                    pass
            self.log(f'已把 {moved} 个旧数据文件迁移到"数据文件"文件夹')

    def _safe_after(self, delay_ms, func, *args):
        """窗口已销毁时 after 会抛异常，采集线程收尾必须不受其影响"""
        try:
            self.root.after(delay_ms, func, *args)
        except Exception:
            pass

    # ---------- 历史导出记录 ----------
    def _init_history(self):
        loaded = None
        corrupt = False
        if os.path.exists(self.history_path):
            try:
                with open(self.history_path, 'r', encoding='utf-8') as f:
                    raw = json.load(f)
                # 必须是字符串列表（dict/数字/嵌套都算损坏），空串直接滤掉
                if isinstance(raw, list) and all(isinstance(x, str) for x in raw):
                    loaded = {x for x in raw if x}
            except Exception:
                pass
            if loaded is None:
                corrupt = True
                try:
                    os.replace(self.history_path, self.history_path + '.corrupt.bak')
                except OSError:
                    pass
        else:
            # 旁边没有历史文件（比如刚把单个EXE发给别人）→ 先尝试内嵌快照
            try:
                with open(self.resource_path(HISTORY_FILE), 'r', encoding='utf-8') as f:
                    raw = json.load(f)
                if isinstance(raw, list) and all(isinstance(x, str) for x in raw):
                    loaded = {x for x in raw if x}
                    self.log(f'从内嵌快照恢复导出记录 {len(loaded)} 个电话')
            except Exception:
                pass

        if loaded is not None:
            self.exported_phones = loaded
            self._save_history()  # 立刻写到旁边，后续读写都走真实文件
            return

        # 文件不存在或已损坏 → 从内置数据重建
        self.exported_phones = set()
        try:
            with open(self.resource_path('data_all.json'), 'r', encoding='utf-8') as f:
                raw = json.load(f)
            # 兼容两种结构：纯列表 [...] 或字典 {'rows': [...]}
            rows_data = raw if isinstance(raw, list) else raw.get('rows', [])
            for rec in rows_data:
                normalize_record(rec)
                for p in split_phones(rec.get('tel')):
                    np_ = normalize_phone(p)
                    if np_:  # 空串不进集合（"暂无"之类洗出来是空）
                        self.exported_phones.add(np_)
        except Exception as e:
            self.log(f'[警告] 内置数据读取失败: {e}')
        self._save_history()

        if corrupt:
            messagebox.showwarning(
                '导出记录已重建',
                '检测到"已导出电话记录.json"文件损坏，已自动从内置数据重建。\n'
                '损坏的原文件备份为 .corrupt.bak。\n\n'
                '注意：如果你近期导出过新采集的数据，那些记录可能丢失，\n'
                '下次"导出新增"时可能出现少量重复，请知悉。')

    def _save_history(self):
        """返回 True/False：去重记录是命根子，写不进盘必须让调用方知道"""
        try:
            tmp = self.history_path + '.tmp'
            with open(tmp, 'w', encoding='utf-8') as f:
                json.dump(sorted(self.exported_phones), f, ensure_ascii=False)
            os.replace(tmp, self.history_path)  # 原子替换，无中间空窗
            try:
                shutil.copyfile(self.history_path, os.path.join(self.state_dir, HISTORY_BAK))
            except OSError:
                pass
            return True
        except Exception:
            return False

    # ---------- 导出日志 ----------
    def _export_log_path(self):
        return os.path.join(self.state_dir, EXPORT_LOG_FILE)

    def _load_export_log(self):
        try:
            with open(self._export_log_path(), 'r', encoding='utf-8') as f:
                raw = json.load(f)
            if isinstance(raw, list):
                return [x for x in raw if isinstance(x, dict)]
        except Exception:
            pass
        return []

    def _append_export_log(self, label, count, phones, files):
        try:
            logs = self._load_export_log()
            logs.append({
                'time': datetime.now().strftime('%Y-%m-%d %H:%M:%S'),
                'type': label,
                'count': count,
                'phones': phones,
                'files': len(files),
                'dir': os.path.dirname(files[0]) if files else '',
            })
            tmp = self._export_log_path() + '.tmp'
            with open(tmp, 'w', encoding='utf-8') as f:
                json.dump(logs, f, ensure_ascii=False, indent=1)
            os.replace(tmp, self._export_log_path())
        except Exception:
            pass

    def _phones_of(self, rec):
        # 惰性缓存：每条记录只算一次正则，结果存在 rec['_phones']
        # （_save_cache 已过滤 _ 前缀字段，不会写进缓存文件）
        ph = rec.get('_phones')
        if ph is None:
            ph = {normalize_phone(p) for p in split_phones(rec.get('tel')) if normalize_phone(p)}
            rec['_phones'] = ph
        return ph

    def _is_exported(self, rec):
        ph = self._phones_of(rec)
        return bool(ph) and ph.issubset(self.exported_phones)

    def _is_new(self, rec):
        ph = self._phones_of(rec)
        return bool(ph) and not ph.issubset(self.exported_phones)

    # ---------- 缓存 ----------
    def _load_cache(self):
        # 内置数据打标，保存缓存时只存非内置的新数据（不依赖下标切分）
        builtin_ok = False
        try:
            with open(self.resource_path('data_all.json'), 'r', encoding='utf-8') as f:
                raw = json.load(f)
            # 兼容两种结构：纯列表 [...] 或字典 {'rows': [...]}
            rows_data = raw if isinstance(raw, list) else raw.get('rows', [])
            for rec in rows_data:
                rec = normalize_record(rec)
                rec['_builtin'] = True
                self.records.append(rec)
            builtin_ok = True
        except Exception as e:
            self.log(f'[警告] 内置数据加载失败: {e}')
        if not builtin_ok or not any(r.get('_builtin') for r in self.records):
            if os.path.exists(self.resource_path('data_all.json')):
                # 数据包文件存在但没读出来（损坏）→ 弹窗提醒修复
                messagebox.showerror(
                    '内置数据加载失败',
                    '内置的商户数据包没能加载进来，软件里只会显示新采集的数据。\n'
                    '请重新下载软件或联系作者修复。')
            # 数据包文件压根不存在（开源用户自己采集的正常路径）→ 静默，不打扰
        # 缓存来源：旁边有真实文件用真实文件；没有则尝试内嵌快照（单EXE分发场景）
        cache_src = None
        if os.path.exists(self.cache_path):
            cache_src = self.cache_path
        else:
            try:
                embedded = self.resource_path(CACHE_FILE)
                if os.path.exists(embedded):
                    cache_src = embedded
            except Exception:
                pass
        if cache_src:
            try:
                with open(cache_src, 'r', encoding='utf-8') as f:
                    raw = json.load(f)
                # 兼容两种结构：纯列表 [...] 或字典 {'rows': [...]}
                cache_rows = raw if isinstance(raw, list) else raw.get('rows', [])
                seen = set()
                for rec in self.records:
                    for p in split_phones(rec.get('tel')):
                        np_ = normalize_phone(p)
                        if np_:
                            seen.add(np_)
                added = 0
                for rec in cache_rows:
                    normalize_record(rec)
                    ph = self._phones_of(rec)
                    # 与导出/展示逻辑统一：全部电话都见过才算重复，部分新号码的记录保留
                    if not ph or ph.issubset(seen):
                        continue
                    self.records.append(rec)
                    seen |= ph
                    added += 1
                if added:
                    self.log(f'加载缓存新增数据 {added} 条')
                if cache_src != self.cache_path:
                    # 读的是内嵌快照（旁边还没有真实缓存文件）→ 立刻落盘，
                    # 保证"数据文件"文件夹完整，整个文件夹拷走不丢新采数据
                    self._save_cache()
            except Exception as e:
                self.log(f'[警告] 采集缓存加载失败: {e}')

    def _save_cache(self):
        """返回 True/False：缓存写不进盘必须让调用方知道"""
        try:
            new_recs = [{k: v for k, v in r.items() if not k.startswith('_')}
                        for r in self.records if not r.get('_builtin')]
            tmp = self.cache_path + '.tmp'
            with open(tmp, 'w', encoding='utf-8') as f:
                json.dump({'headers': HEADERS, 'rows': new_recs,
                           'saved': datetime.now().isoformat(timespec='seconds')},
                          f, ensure_ascii=False)
            os.replace(tmp, self.cache_path)
            return True
        except Exception:
            return False

    # ---------- 地区级联（代际标记防乱序） ----------
    def _init_region_cascade(self):
        self.province_var.set('湖北省')
        self._on_province()

    def _on_province(self, event=None):
        pname = self.province_var.get()
        code_map = {n: c for c, n in PROVINCES}
        pcode = code_map.get(pname)
        if not pcode:
            return
        self.city_combo['values'] = []
        self.city_var.set('加载中...')
        self.district_combo['values'] = [ALL_DISTRICTS]
        self.district_var.set(ALL_DISTRICTS)
        self._district_items = []
        self._refresh_table()  # 市还在路上，先按省过滤刷新，不让用户盯着旧数据

        self._region_gen += 1
        gen = self._region_gen

        def work():
            items = fetch_sub_districts(pcode)
            self._safe_after(0, self._fill_cities, items, gen)
        threading.Thread(target=work, daemon=True).start()

    def _fill_cities(self, items, gen):
        if gen != self._region_gen:
            return  # 过期响应：用户已经又切了省份，丢弃
        if not items:
            self.city_combo['values'] = []
            self.city_var.set('')
            self._refresh_table()  # 市列表拉不到，也先按省过滤刷新
            self.log('[地区] 城市列表获取失败（网络问题），重新选择省份即可重试')
            return
        self._city_items = items
        names = [n for _, n in items]
        self.city_combo['values'] = names
        default = next((n for _, n in items if '宜昌' in n), names[0])
        self.city_var.set(default)
        self._on_city()

    def _on_city(self, event=None):
        cname = self.city_var.get()
        code_map = {n: c for c, n in self._city_items}
        ccode = code_map.get(cname)
        if not ccode:
            return
        self.district_combo['values'] = ['加载中...']
        self.district_var.set('加载中...')
        self._refresh_table()  # 区县还在路上，先按省+市过滤刷新

        self._region_gen += 1
        gen = self._region_gen

        def work():
            items = fetch_sub_districts(ccode)
            if not items and ccode == '420500':  # 宜昌离线兜底
                items = list(YICHANG_DISTRICTS)
            self._safe_after(0, self._fill_districts, items, gen)
        threading.Thread(target=work, daemon=True).start()

    def _fill_districts(self, items, gen):
        if gen != self._region_gen:
            return  # 过期响应，丢弃
        if not items:
            # 该市没有下辖区县（直辖市的区、直筒子市等）→ 用市本级 adcode 采集
            cname = self.city_var.get()
            code_map = {n: c for c, n in self._city_items}
            ccode = code_map.get(cname, '')
            if ccode:
                items = [(ccode, cname)]
                self.log(f'[地区] {cname} 无下辖区县，将按市本级采集')
        self._district_items = items
        names = [n for _, n in items]
        self.district_combo['values'] = [ALL_DISTRICTS] + names
        self.district_var.set(ALL_DISTRICTS)
        self._refresh_table()

    def _selected_districts(self):
        """采集用：全部区县 -> 整个列表；具体区县 -> 单元素列表"""
        if not self._district_items:
            return []
        sel = self.district_var.get()
        if sel in ('', ALL_DISTRICTS, '加载中...'):
            return list(self._district_items)
        code = {n: c for c, n in self._district_items}.get(sel)
        return [(code, sel)] if code else []

    # ---------- 行业级联 ----------
    def _on_ind(self, event=None):
        ind = self.ind_var.get()
        if ind == '全部' or ind not in IND_TREE:
            self.sub_combo['values'] = ['全部']
            self.sub_var.set('全部')
        else:
            self.sub_combo['values'] = list(IND_TREE[ind].keys())
            self.sub_var.set('全部')
        self._refresh_table()

    # ---------- 搜索防抖 ----------
    def _on_search_change(self, *args):
        if self._search_after is not None:
            try:
                self.root.after_cancel(self._search_after)
            except Exception:
                pass
        self._search_after = self.root.after(300, self._refresh_table)

    # ---------- UI ----------
    def _build_ui(self):
        root = self.root
        root.configure(bg='#f5f6fa')

        # 顶部统计（两行：第一行商户/电话总量，第二行当前显示）
        top = tk.Frame(root, bg='#2c3e50', height=64)
        top.pack(fill='x')
        top.pack_propagate(False)
        self.stats_label = tk.Label(top, text='加载中...', font=('微软雅黑', 11, 'bold'),
                                    fg='white', bg='#2c3e50', justify='left', anchor='w')
        self.stats_label.pack(side='left', padx=15)
        tk.Label(top, text='绿色版：数据存于软件旁"数据文件"文件夹，勿删；整个文件夹可拷贝带走',
                 font=('微软雅黑', 9), fg='#95a5a6', bg='#2c3e50').pack(side='right', padx=15)

        # ===== 选择行：省 → 市 → 区县 → 行业 → 细分类目 → 范围 → 搜索 =====
        sel = tk.Frame(root, bg='#f5f6fa')
        sel.pack(fill='x', padx=10, pady=(8, 2))

        tk.Label(sel, text='省:', font=('微软雅黑', 10), bg='#f5f6fa').pack(side='left')
        self.province_var = tk.StringVar()
        self.province_combo = ttk.Combobox(sel, textvariable=self.province_var,
                                           values=[n for _, n in PROVINCES], state='readonly', width=11)
        self.province_combo.pack(side='left', padx=(2, 8))
        self.province_combo.bind('<<ComboboxSelected>>', self._on_province)

        tk.Label(sel, text='市:', font=('微软雅黑', 10), bg='#f5f6fa').pack(side='left')
        self.city_var = tk.StringVar()
        self.city_combo = ttk.Combobox(sel, textvariable=self.city_var,
                                       values=[], state='readonly', width=10)
        self.city_combo.pack(side='left', padx=(2, 8))
        self.city_combo.bind('<<ComboboxSelected>>', self._on_city)

        tk.Label(sel, text='区县:', font=('微软雅黑', 10), bg='#f5f6fa').pack(side='left')
        self.district_var = tk.StringVar(value=ALL_DISTRICTS)
        self.district_combo = ttk.Combobox(sel, textvariable=self.district_var,
                                           values=[ALL_DISTRICTS], state='readonly', width=13)
        self.district_combo.pack(side='left', padx=(2, 8))
        self.district_combo.bind('<<ComboboxSelected>>', lambda e: self._refresh_table())

        tk.Label(sel, text='行业:', font=('微软雅黑', 10), bg='#f5f6fa').pack(side='left')
        self.ind_var = tk.StringVar(value='餐饮服务')
        self.ind_combo = ttk.Combobox(sel, textvariable=self.ind_var,
                                      values=['全部'] + IND_NAMES, state='readonly', width=11)
        self.ind_combo.pack(side='left', padx=(2, 4))
        self.ind_combo.bind('<<ComboboxSelected>>', self._on_ind)

        tk.Label(sel, text='细分类目:', font=('微软雅黑', 10), bg='#f5f6fa').pack(side='left')
        self.sub_var = tk.StringVar(value='全部')
        self.sub_combo = ttk.Combobox(sel, textvariable=self.sub_var,
                                      values=list(IND_TREE['餐饮服务'].keys()), state='readonly', width=13)
        self.sub_combo.pack(side='left', padx=(2, 8))
        self.sub_combo.bind('<<ComboboxSelected>>', lambda e: self._refresh_table())

        tk.Label(sel, text='范围:', font=('微软雅黑', 10), bg='#f5f6fa').pack(side='left')
        self.scope_var = tk.StringVar(value=SCOPE_ALL)
        self.scope_combo = ttk.Combobox(sel, textvariable=self.scope_var,
                                        values=[SCOPE_ALL, SCOPE_NEW, SCOPE_EXPORTED],
                                        state='readonly', width=11)
        self.scope_combo.pack(side='left', padx=(2, 8))
        self.scope_combo.bind('<<ComboboxSelected>>', lambda e: self._refresh_table())

        tk.Label(sel, text='搜索:', font=('微软雅黑', 10), bg='#f5f6fa').pack(side='left')
        self.search_var = tk.StringVar()
        self.search_var.trace_add('write', self._on_search_change)
        tk.Entry(sel, textvariable=self.search_var, font=('微软雅黑', 10), width=14).pack(side='left', padx=(2, 0))

        # ===== 操作行：Key | 采集 | 分割 | 导出 =====
        ops = tk.Frame(root, bg='#f5f6fa')
        ops.pack(fill='x', padx=10, pady=4)

        tk.Label(ops, text='API Key:', font=('微软雅黑', 9), bg='#f5f6fa').pack(side='left')
        self.key_var = tk.StringVar(value=AMAP_KEY)
        tk.Entry(ops, textvariable=self.key_var, font=('微软雅黑', 9), width=24).pack(side='left', padx=(2, 10))

        self.crawl_btn = tk.Button(ops, text='▶ 开始采集', font=('微软雅黑', 10, 'bold'),
                                   bg='#27ae60', fg='white', relief='flat',
                                   padx=14, command=self.toggle_crawl)
        self.crawl_btn.pack(side='left', padx=(0, 12))

        tk.Label(ops, text='分割:', font=('微软雅黑', 10), bg='#f5f6fa').pack(side='left')
        self.split_mode_var = tk.StringVar(value='每份条数')
        ttk.Combobox(ops, textvariable=self.split_mode_var,
                     values=['每份条数', '分成份数'], state='readonly', width=9).pack(side='left', padx=(2, 2))
        self.split_value_var = tk.StringVar(value='200')
        tk.Entry(ops, textvariable=self.split_value_var, font=('微软雅黑', 10), width=6).pack(side='left', padx=(0, 2))
        tk.Label(ops, text='(填1或大于总数=不分割)', font=('微软雅黑', 8), bg='#f5f6fa', fg='#7f8c8d').pack(side='left', padx=(0, 12))

        self.export_view_btn = tk.Button(ops, text='导出当前显示 Excel', font=('微软雅黑', 10),
                                         bg='#2980b9', fg='white', relief='flat',
                                         padx=12, command=lambda: self.export_excel(False))
        self.export_view_btn.pack(side='left', padx=4)
        self.export_new_btn = tk.Button(ops, text='导出新增 Excel(自动去重)', font=('微软雅黑', 10, 'bold'),
                                        bg='#e67e22', fg='white', relief='flat',
                                        padx=12, command=lambda: self.export_excel(True))
        self.export_new_btn.pack(side='left', padx=4)
        tk.Button(ops, text='📁 历史记录', font=('微软雅黑', 10),
                  bg='#8e44ad', fg='white', relief='flat',
                  padx=12, command=self.open_history_window).pack(side='left', padx=(12, 4))

        # 日志行
        self.log_label = tk.Label(root, text='就绪', font=('微软雅黑', 9), bg='#f5f6fa',
                                  fg='#7f8c8d', anchor='w')
        self.log_label.pack(fill='x', padx=12)

        # 表格
        cols = ('name', 'industry', 'sub', 'district', 'address', 'tel', 'location')
        heads = ('商户名称', '行业', '二级类目', '区县', '地址', '联系电话', '经纬度')
        widths = (230, 100, 110, 110, 330, 190, 170)
        anchors = ('w', 'center', 'center', 'center', 'w', 'center', 'center')

        tf = tk.Frame(root)
        tf.pack(fill='both', expand=True, padx=10, pady=(0, 10))
        vsb = ttk.Scrollbar(tf, orient='vertical')
        self.tree = ttk.Treeview(tf, columns=cols, show='headings', yscrollcommand=vsb.set)
        vsb.config(command=self.tree.yview)
        for c, h, w, a in zip(cols, heads, widths, anchors):
            self.tree.heading(c, text=h)
            self.tree.column(c, width=w, anchor=a, minwidth=60)
        self.tree.pack(side='left', fill='both', expand=True)
        vsb.pack(side='right', fill='y')

        style = ttk.Style()
        style.configure('Treeview', font=('微软雅黑', 9), rowheight=24)
        self.tree.tag_configure('new', background='#fff3cd')

        # 底部
        bot = tk.Frame(root, bg='#ecf0f1', height=24)
        bot.pack(fill='x', side='bottom')
        bot.pack_propagate(False)
        tk.Label(bot, text='黄色行 = 新增未导出 ｜ 上面选什么，表格就显示什么、采集就采什么 ｜ 任何导出都记入去重记录',
                 font=('微软雅黑', 8), bg='#ecf0f1', fg='#7f8c8d').pack(side='left', padx=10)

    # ---------- 日志 ----------
    def log(self, msg):
        try:
            self.log_label.config(text=msg)
        except Exception:
            pass

    # ---------- 筛选/刷新 ----------
    def _refresh_table(self):
        self._search_after = None
        kw = self.search_var.get().strip().lower()
        ind = self.ind_var.get()
        sub = self.sub_var.get()
        province = self.province_var.get()
        city = self.city_var.get()
        district = self.district_var.get()
        scope = self.scope_var.get()

        self.filtered = []
        for rec in self.records:
            ok = True
            # 地区三级过滤：省 → 市 → 区县。没采过的地方就必须显示空
            if ok and province and rec.get('pname') != province:
                ok = False
            if ok and city not in ('', '加载中...') and rec.get('cityname') != city:
                ok = False
            if ok and district not in ('', ALL_DISTRICTS, '加载中...') and rec.get('adname') != district:
                ok = False
            if ok and ind != '全部' and rec.get('industry') != ind:
                ok = False
            if ok and sub != '全部' and rec.get('sub') != sub:
                ok = False
            if ok and scope == SCOPE_NEW and not self._is_new(rec):
                ok = False
            if ok and scope == SCOPE_EXPORTED and not self._is_exported(rec):
                ok = False
            if ok and kw:
                hay = (rec.get('name', '') + rec.get('address', '') + rec.get('tel', '')
                       + rec.get('adname', '') + rec.get('industry', '') + rec.get('sub', '')).lower()
                if kw not in hay:
                    ok = False
            if ok:
                self.filtered.append(rec)

        self.tree.delete(*self.tree.get_children())
        for rec in self.filtered:
            tag = ('new',) if self._is_new(rec) else ()
            self.tree.insert('', 'end', values=(
                rec.get('name', ''), rec.get('industry', ''), rec.get('sub', ''),
                rec.get('adname', ''), rec.get('address', ''),
                rec.get('tel', ''), rec.get('location', ''),
            ), tags=tag)
        self._update_stats()

    def _update_stats(self):
        total = len(self.records)
        all_phones = set()
        new_phones = set()
        new_cnt = 0
        no_tel = 0
        for r in self.records:
            ph = self._phones_of(r)
            all_phones |= ph
            if not ph:
                no_tel += 1
            elif self._is_new(r):
                new_cnt += 1
                new_phones |= ph
        exported_cnt = total - new_cnt - no_tel  # 无电话的既不算新增也不算已导出
        # 已导出电话数：以当前数据里的电话为准，跟历史记录取交集
        exp_phone_cnt = len(all_phones & self.exported_phones)
        text = (f'总数据 {total} 条商户 / {len(all_phones)} 个电话 ｜ '
                f'已导出 {exported_cnt} 条 / {exp_phone_cnt} 个电话 ｜ '
                f'新增未导出 {new_cnt} 条 / {len(new_phones)} 个电话\n'
                f'当前显示 {len(self.filtered)} 条（筛选后）')
        if no_tel:
            text += f' ｜ 无电话 {no_tel} 条'
        self.stats_label.config(text=text)

    # ---------- 采集 ----------
    def toggle_crawl(self):
        if self.crawling:
            self.stop_flag = True
            self.log('正在停止...')
            return
        ind = self.ind_var.get()
        if ind == '全部' or ind not in IND_TREE:
            messagebox.showwarning('提示', '采集前请先选一个具体行业\n（全部行业一次跑量太大，容易超Key的日限额）')
            return
        sub = self.sub_var.get()
        queries = IND_TREE.get(ind, {}).get(sub)
        if not queries:
            messagebox.showwarning('提示', '请选择细分类目')
            return
        districts = self._selected_districts()
        if not districts:
            messagebox.showwarning('提示', '区县列表还没加载好，请稍等或重新选择市')
            return

        global AMAP_KEY
        AMAP_KEY = self.key_var.get().strip() or AMAP_KEY
        if not AMAP_KEY:
            messagebox.showwarning('缺少 API Key',
                                   '请先在上方"API Key"框里填入你的高德 Key。\n\n'
                                   '没有 Key？到高德开放平台（lbs.amap.com）注册账号，\n'
                                   '创建"Web服务"类型的应用即可免费获得。')
            return

        # 先用一条请求验证Key有效性 + 额度，失败直接提示，不白跑全部类目
        # 验证期间锁定下拉框：防止用户验证完前改了选项，采的和显示的不一致
        self.crawl_btn.config(state='disabled', text='验证中...')
        for cb in (self.province_combo, self.city_combo, self.district_combo,
                   self.ind_combo, self.sub_combo):
            cb.config(state='disabled')
        self.log('正在验证 API Key 和剩余额度...')
        first_adcode = districts[0][0]
        first_query = queries[0]

        def check():
            try:
                test = search_poi(first_adcode, first_query, 1)
                if isinstance(test, dict) and test.get('status') == '1':
                    err = None
                else:
                    err = classify_api_error(test)
            except Exception as e:
                err = f'net:{e}'
            self._safe_after(0, self._key_checked, err, districts, ind, queries)
        threading.Thread(target=check, daemon=True).start()

    def _key_checked(self, err, districts, ind, queries):
        if err is not None:
            self.crawl_btn.config(state='normal', text='▶ 开始采集')
            for cb in (self.province_combo, self.city_combo, self.district_combo,
                       self.ind_combo, self.sub_combo):
                cb.config(state='readonly')  # 验证失败：恢复下拉框
            if err == 'daily':
                msg = '今日 API 调用额度已用完。\n\n高德免费Key每天有调用上限，明天额度重置后再采，或换个Key。'
            elif err == 'key':
                msg = 'API Key 无效。\n\n请检查 Key 是否填错（多空格/少字符），或到高德开放平台确认Key状态。'
            elif isinstance(err, str) and err.startswith('net:'):
                msg = f'无法连接高德服务器：\n{err[4:]}\n\n请检查网络后重试。'
            else:
                msg = '高德API拒绝了请求，请检查 Key 和网络后重试。'
            self.log('验证失败：' + msg.split('\n')[0])
            messagebox.showerror('无法开始采集', msg)
            return

        # 验证通过，正式启动采集
        self.crawling = True
        self.stop_flag = False
        self.crawl_btn.config(state='normal', text='■ 停止采集', bg='#e74c3c')
        self.export_view_btn.config(state='disabled')
        self.export_new_btn.config(state='disabled')
        for cb in (self.province_combo, self.city_combo, self.district_combo,
                   self.ind_combo, self.sub_combo):
            cb.config(state='disabled')
        threading.Thread(target=self._crawl_worker,
                         args=(districts, ind, queries), daemon=True).start()

    def _crawl_worker(self, districts, ind_name, queries):
        added = 0
        skipped = 0
        snapshot = list(self.records)  # 先快照：中途任何异常也能保住已采集部分
        try:
            existing = set(self.exported_phones)
            for rec in self.records:
                for p in split_phones(rec.get('tel')):
                    np_ = normalize_phone(p)
                    if np_:
                        existing.add(np_)

            total_steps = len(districts) * len(queries)
            step = 0
            last_refresh = time.time()
            daily_out = False

            for adcode, dname in districts:
                if self.stop_flag:
                    break
                for kind, value, label in queries:
                    if self.stop_flag:
                        break
                    step += 1
                    self._safe_after(0, lambda s=step, t=total_steps, d=dname, l=label:
                                     self.log(f'[{s}/{t}] {d} · {l} ... 新增{added} 跳过{skipped}'))
                    page = 1
                    qps_retries = 0
                    while page <= 40:
                        if self.stop_flag:
                            break
                        try:
                            data = search_poi(adcode, (kind, value, label), page)
                        except Exception as e:
                            self._safe_after(0, lambda m=str(e): self.log(f'[网络错误] {m}，已跳过'))
                            break
                        if not isinstance(data, dict) or data.get('status') != '1':
                            err = classify_api_error(data)
                            if err == 'daily':
                                daily_out = True
                                self.stop_flag = True
                                self._safe_after(0, lambda: self.log('[额度] 今日额度已用完，提前结束'))
                                break
                            if err == 'qps' and qps_retries < 3:
                                qps_retries += 1
                                time.sleep(1.5)
                                continue  # 重试本页
                            break  # 其他错误：跳过该类目
                        qps_retries = 0
                        pois = data.get('pois') or []
                        if not pois:
                            break
                        for poi in pois:
                            rec = {
                                'name': safe_str(poi.get('name')),
                                'pname': safe_str(poi.get('pname')),
                                'cityname': safe_str(poi.get('cityname')),
                                'adname': safe_str(poi.get('adname')),
                                'address': safe_str(poi.get('address')),
                                'tel': safe_str(poi.get('tel')),
                                'location': safe_str(poi.get('location')),
                                'type': safe_str(poi.get('type')),
                                'industry': ind_name,
                                'sub': label,
                            }
                            phones = {normalize_phone(p) for p in split_phones(rec['tel'])}
                            phones.discard('')
                            rec['_phones'] = phones  # 顺手缓存，主线程统计不再重算
                            if not phones:
                                skipped += 1
                                continue
                            # 与展示/导出逻辑统一：全部电话都已存在才算重复（issubset）
                            # 含新号码的商户保留——避免"旧号码+新号码"的商户整条丢失
                            if phones.issubset(existing):
                                skipped += 1
                                continue
                            existing |= phones
                            snapshot.append(rec)
                            added += 1
                        if len(pois) < 25:
                            break
                        page += 1
                        time.sleep(0.3)  # 免费Key QPS上限约3次/秒
                        now = time.time()
                        if now - last_refresh > 3:
                            last_refresh = now
                            self.records = list(snapshot)
                            self._safe_after(0, self._refresh_table)
                if daily_out:
                    break
        except Exception as e:
            # 采集中途出任何意外：已采集部分照常保存，界面照常恢复
            self._safe_after(0, lambda m=str(e): self.log(f'[采集异常中断] {m}，已采集部分已保存'))
        finally:
            # 无论正常结束/手动停止/额度尽/异常，都保存已采集的数据
            self.records = list(snapshot)
            save_ok = self._save_cache()
            # 保存完成才置位：关窗等待逻辑盯着这个标志，不能保存到一半放行
            self.crawling = False
        self._safe_after(0, self._crawl_done, added, skipped, save_ok)

    def _crawl_done(self, added, skipped, save_ok=True):
        self.crawl_btn.config(state='normal', text='▶ 开始采集', bg='#27ae60')
        self.export_view_btn.config(state='normal')
        self.export_new_btn.config(state='normal')
        for cb in (self.province_combo, self.city_combo, self.district_combo,
                   self.ind_combo, self.sub_combo):
            cb.config(state='readonly')
        # 采集结束：历史窗口如果开着，恢复两个清空按钮
        try:
            if self._hist_win is not None and self._hist_win.winfo_exists():
                self._hist_clear_history_btn.config(state='normal')
                self._hist_clear_cache_btn.config(state='normal')
        except Exception:
            pass
        self._refresh_table()
        if not save_ok:
            messagebox.showerror('采集数据保存失败',
                                 f'这次采到的 {added} 条数据没能写入数据文件！\n\n'
                                 f'保存位置：{self.cache_path}\n\n'
                                 f'请检查：U盘是否拔出、磁盘是否已满、文件夹是否被占用。\n'
                                 f'数据目前还在内存里，修复后随便再采集一条即可触发重新保存。')
        self.log(f'采集完成：新增 {added} 条（自动跳过已有/已导出 {skipped} 条），点"导出新增"即可拿到这批新数据')
        messagebox.showinfo('采集完成',
                            f'新增 {added} 条全新商户（已跳过 {skipped} 条重复/无电话）\n\n'
                            f'新增行已黄色标出，点"导出新增 Excel"即可导出这批数据。')

    # ---------- 关窗：采集中途也要把数据存下来 ----------
    def _on_close(self):
        if self.crawling:
            self.stop_flag = True
            self.log('正在保存已采集数据，请稍候...')
            self._wait_worker_exit(0)
        else:
            self.root.destroy()

    def _wait_worker_exit(self, waited):
        # worker 收尾时会 _save_cache 并把 crawling 置 False；最多等20秒兜底
        if not self.crawling or waited >= 200:
            self.root.destroy()
            return
        self.root.after(100, lambda: self._wait_worker_exit(waited + 1))

    # ---------- 历史记录管理窗口 ----------
    def open_history_window(self):
        # 防多开：已开着就拉到最前，不再开新窗口（多开会导致控件引用互相覆盖）
        if self._hist_win is not None:
            try:
                if self._hist_win.winfo_exists():
                    self._hist_win.lift()
                    self._hist_win.focus_force()
                    return
            except Exception:
                pass
        win = tk.Toplevel(self.root)
        self._hist_win = win
        win.title('历史记录管理')
        win.geometry('900x620')
        win.configure(bg='#f5f6fa')
        win.transient(self.root)

        # 状态摘要
        summary = tk.Frame(win, bg='white', relief='solid', bd=1)
        summary.pack(fill='x', padx=12, pady=(12, 6))
        self._hist_summary_var = tk.StringVar()
        tk.Label(summary, textvariable=self._hist_summary_var, font=('微软雅黑', 10),
                 bg='white', justify='left', anchor='w').pack(fill='x', padx=12, pady=(10, 2))
        if self._state_dir_fallback:
            loc_text = f'⚠ 软件目录不可写，数据临时保存在：{self.state_dir}'
        else:
            loc_text = f'数据文件在软件旁边的"数据文件"文件夹里（勿删，整个文件夹拷走即用）：{self.state_dir}'
        tk.Label(summary, text=loc_text,
                 font=('微软雅黑', 8), fg='#95a5a6', bg='white', anchor='w').pack(fill='x', padx=12, pady=(0, 8))

        # 两个标签页：已导出电话清单 / 导出操作日志
        nb = ttk.Notebook(win)
        nb.pack(fill='both', expand=True, padx=12, pady=(6, 0))

        # —— 标签1：已导出电话清单（以往历史记录）——
        tab1 = tk.Frame(nb, bg='#f5f6fa')
        nb.add(tab1, text=' 已导出电话清单 ')
        sf = tk.Frame(tab1, bg='#f5f6fa')
        sf.pack(fill='x', pady=6)
        tk.Label(sf, text='搜索:', font=('微软雅黑', 9), bg='#f5f6fa').pack(side='left')
        self._phone_filter_var = tk.StringVar()
        e = tk.Entry(sf, textvariable=self._phone_filter_var, font=('微软雅黑', 9), width=24)
        e.pack(side='left', padx=(4, 8))
        e.bind('<KeyRelease>', lambda _e: self._fill_phone_list())
        self._phone_count_var = tk.StringVar()
        tk.Label(sf, textvariable=self._phone_count_var, font=('微软雅黑', 9),
                 bg='#f5f6fa', fg='#7f8c8d').pack(side='left')
        pf = tk.Frame(tab1)
        pf.pack(fill='both', expand=True)
        psb = ttk.Scrollbar(pf, orient='vertical')
        self._phone_tree = ttk.Treeview(pf, show='headings', yscrollcommand=psb.set,
                                        columns=('idx', 'phone'))
        psb.config(command=self._phone_tree.yview)
        self._phone_tree.heading('idx', text='#')
        self._phone_tree.heading('phone', text='已导出过的电话')
        self._phone_tree.column('idx', width=60, anchor='center')
        self._phone_tree.column('phone', width=260, anchor='w')
        self._phone_tree.pack(side='left', fill='both', expand=True)
        psb.pack(side='right', fill='y')

        # —— 标签2：导出操作日志 ——
        tab2 = tk.Frame(nb, bg='#f5f6fa')
        nb.add(tab2, text=' 导出操作日志 ')
        tk.Label(tab2, text='每次导出的时间、条数、文件数（新版起自动记录）：',
                 font=('微软雅黑', 9), bg='#f5f6fa', anchor='w', fg='#7f8c8d').pack(fill='x', pady=(6, 2))
        lf = tk.Frame(tab2)
        lf.pack(fill='both', expand=True)
        lsb = ttk.Scrollbar(lf, orient='vertical')
        self._log_tree = ttk.Treeview(lf, show='headings', yscrollcommand=lsb.set,
                                      columns=('time', 'type', 'count', 'phones', 'files', 'dir'))
        lsb.config(command=self._log_tree.yview)
        for c, h, w in (('time', '时间', 150), ('type', '类型', 90), ('count', '条数', 70),
                        ('phones', '新增电话', 80), ('files', '文件数', 60), ('dir', '保存目录', 300)):
            self._log_tree.heading(c, text=h)
            self._log_tree.column(c, width=w, anchor='center' if c != 'dir' else 'w')
        self._log_tree.pack(side='left', fill='both', expand=True)
        lsb.pack(side='right', fill='y')

        # 按钮行
        bf = tk.Frame(win, bg='#f5f6fa')
        bf.pack(fill='x', padx=12, pady=10)
        self._hist_clear_history_btn = tk.Button(
            bf, text='清空导出历史', font=('微软雅黑', 10, 'bold'),
            bg='#c0392b', fg='white', relief='flat', padx=14,
            command=lambda: self._clear_history(win))
        self._hist_clear_history_btn.pack(side='left', padx=(0, 8))
        self._hist_clear_cache_btn = tk.Button(
            bf, text='清空采集缓存', font=('微软雅黑', 10, 'bold'),
            bg='#d35400', fg='white', relief='flat', padx=14,
            command=lambda: self._clear_cache(win))
        self._hist_clear_cache_btn.pack(side='left', padx=8)
        if self.crawling:
            # 采集中禁止清空：worker收尾会按快照回写，清空会被覆盖失效
            self._hist_clear_history_btn.config(state='disabled')
            self._hist_clear_cache_btn.config(state='disabled')
        tk.Button(bf, text='刷新', font=('微软雅黑', 10), relief='flat', padx=14,
                  command=self._reload_history_window).pack(side='left', padx=8)
        tk.Button(bf, text='关闭', font=('微软雅黑', 10), relief='flat', padx=14,
                  command=win.destroy).pack(side='right')

        self._reload_history_window()

    def _fill_phone_list(self):
        """填充已导出电话清单（带搜索过滤）"""
        for item in self._phone_tree.get_children():
            self._phone_tree.delete(item)
        kw = self._phone_filter_var.get().strip()
        phones = sorted(self.exported_phones)
        shown = 0
        for p in phones:
            if kw and kw not in p:
                continue
            shown += 1
            self._phone_tree.insert('', 'end', values=(shown, p))
        self._phone_count_var.set(f'共 {len(phones)} 个' + (f'，匹配 {shown} 个' if kw else ''))

    def _reload_history_window(self):
        """刷新历史窗口的电话清单、日志列表和摘要"""
        self._fill_phone_list()
        for item in self._log_tree.get_children():
            self._log_tree.delete(item)
        for log in reversed(self._load_export_log()):  # 最新在前
            self._log_tree.insert('', 'end', values=(
                log.get('time', ''), log.get('type', ''), log.get('count', 0),
                log.get('phones', 0), log.get('files', 0), log.get('dir', '')))
        cache_cnt = sum(1 for r in self.records if not r.get('_builtin'))
        self._hist_summary_var.set(
            f'已导出电话记录：{len(self.exported_phones)} 个电话（用于"导出新增"时自动去重）\n'
            f'采集缓存：{cache_cnt} 条新采集的商户数据（内置数据包不算在内）')

    def _clear_history(self, win):
        if self.crawling:
            messagebox.showwarning('采集中', '正在采集数据，清空操作会被采集结果覆盖。\n请先停止采集再清空。', parent=win)
            return
        if not messagebox.askyesno(
                '确认清空导出历史',
                f'将清空全部 {len(self.exported_phones)} 个已导出电话记录。\n\n'
                f'清空后：所有商户都会变成"新增未导出"（黄色），\n'
                f'下次"导出新增"会导出全量数据，可能和以前的导出重复。\n\n'
                f'确定清空吗？（原文件会备份为 .bak）', parent=win):
            return
        try:
            if os.path.exists(self.history_path):
                shutil.copyfile(self.history_path, os.path.join(self.state_dir, HISTORY_BAK))
        except OSError:
            pass
        self.exported_phones = set()
        if self._save_history():  # 写空列表覆盖（不是删文件——删了会从内嵌快照复活）
            self._refresh_table()
            self._reload_history_window()
            self.log('已清空导出历史记录（备份在 .bak）')
            messagebox.showinfo('已清空', '导出历史已清空，原记录备份为 已导出电话记录.json.bak。\n'
                                         '现在所有数据都显示为"新增未导出"。', parent=win)
        else:
            messagebox.showerror('清空失败',
                                 f'去重记录文件写入失败，清空未生效。\n保存位置：{self.state_dir}\n'
                                 f'请检查磁盘/占用后重试。', parent=win)

    def _clear_cache(self, win):
        if self.crawling:
            messagebox.showwarning('采集中', '正在采集数据，清空操作会被采集结果覆盖。\n请先停止采集再清空。', parent=win)
            return
        cache_cnt = sum(1 for r in self.records if not r.get('_builtin'))
        if not messagebox.askyesno(
                '确认清空采集缓存',
                f'将删除全部 {cache_cnt} 条新采集的数据，只保留内置数据包。\n\n'
                f'注意：这些商户的电话如果已在导出历史里，去重记录仍然保留。\n\n'
                f'确定清空吗？', parent=win):
            return
        self.records = [r for r in self.records if r.get('_builtin')]
        if self._save_cache():  # 写空缓存文件覆盖（同上，防内嵌快照复活）
            self._refresh_table()
            self._reload_history_window()
            self.log(f'已清空采集缓存（删除 {cache_cnt} 条新采集数据）')
            messagebox.showinfo('已清空', f'已删除 {cache_cnt} 条新采集数据，现在只剩内置数据包。', parent=win)
        else:
            messagebox.showerror('清空失败',
                                 f'缓存文件写入失败，清空未生效。\n保存位置：{self.state_dir}\n'
                                 f'请检查磁盘/占用后重试。', parent=win)

    # ---------- 导出 ----------
    def export_excel(self, only_new):
        if only_new:
            data = [r for r in self.records if self._is_new(r)]
            if not data:
                messagebox.showinfo('提示',
                                    '当前没有"新增未导出"的数据。\n\n'
                                    '表格里所有电话以前都导出过了。\n'
                                    '先点"开始采集"获取新商户，再导出新增。')
                return
        else:
            data = list(self.filtered)
            if not data:
                messagebox.showinfo('提示', '当前筛选结果为空，没有可导出的数据。')
                return

        mode = self.split_mode_var.get()
        try:
            value = int(self.split_value_var.get().strip())
        except ValueError:
            messagebox.showwarning('提示', '分割数量请填整数')
            return
        chunks = make_chunks(data, mode, value)

        label = '新增' if only_new else '当前显示'
        if len(chunks) > 1:
            sizes = '、'.join(str(len(c)) for c in chunks[:6])
            if len(chunks) > 6:
                sizes += ' ...'
            ok = messagebox.askokcancel('确认分割导出',
                                        f'共 {len(data)} 条，将生成 {len(chunks)} 个 Excel 文件：\n'
                                        f'每个文件条数：{sizes}\n\n确定继续吗？')
            if not ok:
                return

        city = self.city_var.get().strip()
        if not city or city in ('加载中...',):
            city = ''
        ts = datetime.now().strftime('%Y%m%d_%H%M%S')
        first = filedialog.asksaveasfilename(
            title='选择导出位置（多文件时自动加序号）',
            defaultextension='.xlsx',
            initialfile=f'{city}商户_{label}_{ts}.xlsx',
            filetypes=[('Excel文件', '*.xlsx')],
        )
        if not first:
            return

        # 多文件模式：目标文件若已存在会被静默覆盖（文件对话框只确认主文件名）
        # → 导出前先检查一遍，给用户换名字的机会
        if len(chunks) > 1:
            base, ext = os.path.splitext(first)
            existed = [os.path.basename(f'{base}_第{i:02d}部分{ext}')
                       for i in range(1, len(chunks) + 1)
                       if os.path.exists(f'{base}_第{i:02d}部分{ext}')]
            if existed:
                show = existed[:5]
                if not messagebox.askyesno(
                        '文件已存在',
                        f'以下 {len(existed)} 个文件已存在，继续导出会被覆盖：\n'
                        + '\n'.join(show) + ('\n...' if len(existed) > 5 else '') +
                        '\n\n确定覆盖吗？（点"否"可以回去换个文件名）'):
                    return

        written = []
        try:
            if len(chunks) == 1:
                write_excel(first, export_rows(chunks[0]))
                written.append(first)
            else:
                base, ext = os.path.splitext(first)
                for i, chunk in enumerate(chunks, 1):
                    path = f'{base}_第{i:02d}部分{ext}'
                    write_excel(path, export_rows(chunk))
                    written.append(path)
        except Exception as e:
            messagebox.showerror('导出中断',
                                 f'写入第 {len(written)+1} 个文件时出错：\n{e}\n\n'
                                 f'已成功 {len(written)} 个。\n'
                                 f'注意：本次导出未计入去重记录，这批数据仍是"新增"状态，'
                                 f'可直接重新导出。')
            return

        new_marks = 0
        for rec in data:
            for p in split_phones(rec.get('tel')):
                np_ = normalize_phone(p)
                if np_ and np_ not in self.exported_phones:
                    self.exported_phones.add(np_)
                    new_marks += 1
        history_saved = self._save_history()
        self._append_export_log(label, len(data), new_marks, written)
        self._refresh_table()

        folder = os.path.dirname(first)
        if len(written) == 1:
            msg = f'已导出 1 个文件（{len(data)} 条）：\n{written[0]}'
        else:
            msg = f'已导出 {len(written)} 个文件，共 {len(data)} 条：\n{folder}\n\n' + \
                  '\n'.join(os.path.basename(p) for p in written[:5]) + \
                  ('\n...' if len(written) > 5 else '')
        if history_saved:
            msg += '\n\n这批商户已全部记入去重记录，下次"导出新增"不会再包含它们。'
        else:
            msg += ('\n\n⚠ 注意：去重记录文件写入失败（U盘拔出/磁盘满/文件夹被占用？）\n'
                    f'保存位置：{self.state_dir}\n'
                    '这批电话目前只在内存里记着，关闭软件后下次导出可能重复！\n'
                    '修复后请再点一次任意导出，即可把记录写回。')
        self.log(f'导出完成：{len(written)} 个文件 / {len(data)} 条'
                 + ('' if history_saved else '（去重记录保存失败！）'))
        messagebox.showinfo('导出完成', msg) if history_saved else messagebox.showwarning('导出完成（有警告）', msg)
        try:
            os.startfile(folder)  # Windows 原生打开文件夹，中文路径稳定
        except Exception:
            try:
                webbrowser.open(f'file:///{folder.replace(os.sep, "/")}')
            except Exception:
                pass

    def run(self):
        self.root.mainloop()


if __name__ == '__main__':
    MerchantViewer().run()
